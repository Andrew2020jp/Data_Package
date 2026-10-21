"""Check whether the creation commit of each L1/L2 record carries an explicit AI signature.

Scope: the 71 records in sheet "Non-L3 Creation Evidence" of data_package/data/adr_dataset.xlsx
(51 L2, 20 L1). Each record has a traced creation commit and a selected (reviewed) commit.
This script only asks: does the *creation* commit carry an explicit AI signature under the
same rule that confirmed the reviewed commits?

Signature rule (paper Sec. II-B; patterns copied from others/verify_records_20260910.py,
the script that confirmed the current records):
  1. a Co-authored-by trailer naming an AI tool or agent
  2. an author or committer identity naming an AI tool or agent
  3. a "generated with/by" marker naming an AI tool
A tool name elsewhere in the message, a generic [bot] identity, or AI wording in the ADR
text is not a signature. Hits matched only by names outside the screening list
(repro/rerun_screening_from_849.py) are kept for human review, not counted as confirmed.

Results are automatic candidates. Human review status is recorded separately and starts
as "not reviewed".

Usage:
  GITHUB_TOKEN=... python scripts/check_non_l3_creation_ai.py          # run / resume
  python scripts/check_non_l3_creation_ai.py --offline                 # cache only
  python scripts/check_non_l3_creation_ai.py --selftest                # rule tests only
The token is read from GITHUB_TOKEN or GH_TOKEN and is never written anywhere.
"""
import csv, json, os, re, sys, time, urllib.error, urllib.parse, urllib.request
from collections import Counter
from pathlib import Path
import openpyxl

PKG = Path(__file__).resolve().parents[1]                 # data_package/
ROOT = PKG.parent                                          # ~/adr
SRC = PKG / 'data/adr_dataset.xlsx'
OUT = PKG / 'analysis/non_l3_creation_ai'
CACHE = OUT / 'api_cache'

# Earlier caches and lifecycle tables, used only as clues (checked by repo, path and SHA).
OLD_CACHE_DIRS = [ROOT / 'others/audits/all_levels_timing_audit_20260906/api',
                  ROOT / 'others/audits/manual_review_level_reaudit_20260910/fresh_github_cache']
OLD_TRACE_CSV = ROOT / 'others/audits/non_l3_creation_20260910/non_l3_creation_evidence.csv'
OLD_SEP02 = (ROOT / 'others/adr_ai_traces_list_english_main.xlsx', 'L2 AI Timing Evidence')
OLD_JUL31 = (ROOT / 'others/archive/sources/original.xlsx', '0731 task lifecycle')

# ------------------------------------------------------------------ rules --
AI_PAT = re.compile(r'anthropic|claude|copilot|cursor|openai|codex|chatgpt|gemini|devin|aider|windsurf|jules|coderabbit|amp\b',
                    re.I)
SCREEN_PAT = re.compile(r'(claude|anthropic|copilot|cursor|gemini|chatgpt|openai|codex|devin|aider|windsurf)', re.I)


def signatures(c):
    """All signature hits on a commit object: list of (kind, text, tool, screening_list_match)."""
    msg = (c.get('commit') or {}).get('message') or ''
    hits = []
    for line in msg.splitlines():
        if re.match(r'\s*co-authored-by\s*:', line, re.I) and AI_PAT.search(line):
            hits.append(('trailer', line.strip()))
    for who in ('author', 'committer'):
        d = (c.get('commit') or {}).get(who) or {}
        blob = f"{d.get('name', '')} <{d.get('email', '')}>"
        if AI_PAT.search(blob):
            hits.append((who, blob))
    for line in msg.splitlines():
        if re.search(r'generated (with|by)', line, re.I) and AI_PAT.search(line):
            hits.append(('generated-with', line.strip()))
    return [(k, t[:200], AI_PAT.search(t).group(0).lower(), bool(SCREEN_PAT.search(t))) for k, t in hits]


def mentions_only(c):
    """Tool name appears in the message but no signature: recorded, not counted."""
    msg = (c.get('commit') or {}).get('message') or ''
    return sorted({m.lower() for m in AI_PAT.findall(msg)}) if not signatures(c) else []


def judge(hits):
    """Automatic judgment for one commit from its signature hits."""
    if not hits:
        return 'no qualifying AI attribution found', ''
    if any(h[3] for h in hits):
        return 'AI attribution (automatic candidate)', ''
    return 'needs human review', 'matched only names outside the screening list (' + \
        ', '.join(sorted({h[2] for h in hits})) + ')'


# ------------------------------------------------------------------ GitHub --
def token():
    return os.environ.get('GITHUB_TOKEN') or os.environ.get('GH_TOKEN')


def cache_file(repo, sha, page):
    return CACHE / f"{repo.replace('/', '__')}__{sha}__p{page}.json"


def old_cached(sha):
    """Commit object with this SHA from an earlier cache, if any (page 1 only)."""
    for d in OLD_CACHE_DIRS:
        if not d.exists():
            continue
        for f in d.glob('*.json'):
            t = f.read_text(errors='ignore')
            if sha not in t:
                continue
            try:
                x = json.loads(t)
            except json.JSONDecodeError:
                continue
            body = x.get('body', x) if isinstance(x, dict) else None
            if isinstance(body, dict) and body.get('sha') == sha and 'commit' in body and 'files' in body:
                return body, str(f.relative_to(ROOT))
    return None, None


def fetch(repo, sha, page, offline):
    f = cache_file(repo, sha, page)
    if f.exists():
        x = json.loads(f.read_text())
        return x['body'], x['source']
    if page == 1:
        body, src = old_cached(sha)
        if body is not None and len(body.get('files') or []) < 100:
            f.write_text(json.dumps({'source': src, 'body': body}))
            return body, src
    if offline:
        return {'_error': 'not cached (offline)'}, ''
    url = f'https://api.github.com/repos/{repo}/commits/{sha}?per_page=100&page={page}'
    req = urllib.request.Request(url, headers={'Accept': 'application/vnd.github+json',
                                               'User-Agent': 'adr-non-l3-check'})
    if token():
        req.add_header('Authorization', 'Bearer ' + token())
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            body = json.loads(r.read())
    except urllib.error.HTTPError as e:
        return {'_error': f'HTTP {e.code}'}, ''          # errors are not cached, so a rerun retries
    except Exception as e:
        return {'_error': type(e).__name__}, ''
    src = f'GitHub API {url.split("?")[0]} (fetched {time.strftime("%Y-%m-%d", time.gmtime())})'
    f.write_text(json.dumps({'source': src, 'body': body}))
    return body, src


def history(repo, path, sha, offline):
    """Commits touching `path` reachable from `sha`, newest first (no rename following)."""
    out = []
    for page in range(1, 11):
        key = CACHE / f"hist__{repo.replace('/', '__')}__{re.sub(r'[^0-9A-Za-z]+', '_', path)}__{sha[:12]}__p{page}.json"
        if key.exists():
            body = json.loads(key.read_text())['body']
        elif offline:
            return None
        else:
            url = (f'https://api.github.com/repos/{repo}/commits?path={urllib.parse.quote(path)}'
                   f'&sha={sha}&per_page=100&page={page}')
            req = urllib.request.Request(url, headers={'Accept': 'application/vnd.github+json',
                                                       'User-Agent': 'adr-non-l3-check'})
            if token():
                req.add_header('Authorization', 'Bearer ' + token())
            try:
                with urllib.request.urlopen(req, timeout=30) as r:
                    body = json.loads(r.read())
            except Exception:
                return None
            key.write_text(json.dumps({'source': url.split('&page')[0], 'body': body}))
        out += body
        if len(body) < 100:
            return out
    return out


def blob(repo, path, ref, offline):
    """Blob SHA of `path` at `ref`, or None."""
    key = CACHE / f"blob__{repo.replace('/', '__')}__{re.sub(r'[^0-9A-Za-z]+', '_', path)}__{ref[:12]}.json"
    if key.exists():
        return json.loads(key.read_text())['body'].get('sha')
    if offline:
        return None
    url = f'https://api.github.com/repos/{repo}/contents/{urllib.parse.quote(path)}?ref={ref}'
    req = urllib.request.Request(url, headers={'Accept': 'application/vnd.github+json', 'User-Agent': 'adr-non-l3-check'})
    if token():
        req.add_header('Authorization', 'Bearer ' + token())
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            body = json.loads(r.read())
    except Exception:
        return None
    body = {'sha': body.get('sha'), 'path': body.get('path')}
    key.write_text(json.dumps({'source': url, 'body': body}))
    return body['sha']


def linked_by_blob(repo, path, old_path, rev, offline):
    """For commits whose file list is truncated: the commit that introduced `path` holds the same
    blob as `old_path` in its parent, i.e. a pure move. Returns a description or None."""
    h = history(repo, path, rev, offline)
    if not h:
        return None
    first = h[-1]['sha']
    c, _ = commit(repo, first, offline)
    if not c.get('parents'):
        return None
    parent = c['parents'][0]['sha']
    new, old = blob(repo, path, first, offline), blob(repo, old_path, parent, offline)
    if new and new == old:
        return f'{old_path} -> {path} in {first[:12]} (same blob {new[:12]}; commit file list truncated)'
    return None


def rename_chain(repo, path, rev, offline):
    """Follow renames back from the reviewed commit. Returns (adding commit SHA, path it added, chain) or None."""
    chain, sha = [], rev
    for _ in range(6):
        h = history(repo, path, sha, offline)
        if not h:
            return None
        first = h[-1]['sha']
        c, _ = commit(repo, first, offline)
        f = next((x for x in c.get('files') or [] if x.get('filename') == path), None)
        if f is None:
            return None
        if f.get('status') == 'added':
            return first, path, chain
        if f.get('status') != 'renamed' or not f.get('previous_filename'):
            return None
        chain.append(f"{f['previous_filename']} -> {path} in {first[:12]}")
        path, sha = f['previous_filename'], first
    return None


def commit(repo, sha, offline):
    """Commit object with all file pages merged."""
    body, src = fetch(repo, sha, 1, offline)
    if '_error' in body:
        return body, src
    files, page = list(body.get('files') or []), 1
    while len(files) == 100 * page and page < 30:
        page += 1
        more, _ = fetch(repo, sha, page, offline)
        if '_error' in more:
            body = dict(body, _files_incomplete=more['_error'])
            break
        files += more.get('files') or []
    return dict(body, files=files), src


# ------------------------------------------------------------------ inputs --
def full_sha(cell):
    m = re.search(r'/commit/([0-9a-f]{40})', (cell.hyperlink.target if cell.hyperlink else '') or '')
    return m.group(1) if m else None


def records():
    ws = openpyxl.load_workbook(SRC)['Non-L3 Creation Evidence']
    head = [c.value for c in next(r for r in ws.iter_rows(max_row=10) if r[0].value == 'Manual Review row')]
    col = {h: i for i, h in enumerate(head)}
    out = []
    for r in ws.iter_rows(min_row=4):
        if not isinstance(r[0].value, int):
            continue
        out.append(dict(row=r[0].value, repo=r[col['Repository']].value, level=r[col['Commit Level']].value,
                        path=r[col['ADR file path']].value,
                        created=r[col['ADR created time (UTC)']].value, reviewed=r[col['Reviewed commit time (UTC)']].value,
                        cre=full_sha(r[col['Creation commit']]), rev=full_sha(r[col['Reviewed commit']]),
                        note=r[col['Notes']].value or ''))
    return out


def table(path, sheet, key):
    if not path.exists():
        return []
    rows = list(openpyxl.load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and key in r)
    return [dict(zip(rows[h], r)) for r in rows[h + 1:] if r and any(r)]


def old_clues():
    trace = {}
    if OLD_TRACE_CSV.exists():
        for r in csv.DictReader(open(OLD_TRACE_CSV, encoding='utf-8-sig')):
            trace[(r['repository'], r['current_adr_path'])] = r
    sep = {(t['Repository'], t['ADR file path']): t for t in table(*OLD_SEP02, 'Lifecycle classification')}
    jul = {(t['Repository'], t['ADR file path']): t for t in table(*OLD_JUL31, 'Lifecycle tag')}
    return trace, sep, jul


# ------------------------------------------------------------------ checks --
def file_check(c, rec, trace, offline):
    """Does the creation commit add this ADR? Returns (status, matched path, result)."""
    files = c.get('files') or []
    paths = [rec['path']]
    t = trace.get((rec['repo'], rec['path']))
    if t and t.get('creation_path_traced') and rec['cre'].startswith(t.get('creation_commit_sha') or '-'):
        paths.append(t['creation_path_traced'])
    for p in paths:
        for f in files:
            if f.get('filename') == p:
                if f.get('status') != 'added':
                    return f.get('status'), p, f'present as {f.get("status")}, not added'
                if p == rec['path']:
                    return 'added', p, 'added at current path'
                # The earlier path comes from an old trace table: check its link to the current path.
                rc = rename_chain(rec['repo'], rec['path'], rec['rev'], offline)
                if rc and rc[0] == rec['cre'] and rc[1] == p:
                    return 'added', p, 'added at earlier path (rename chain: ' + '; '.join(rc[2]) + ')'
                lb = linked_by_blob(rec['repo'], rec['path'], p, rec['rev'], offline)
                if lb:
                    return 'added', p, 'added at earlier path (' + lb + ')'
                return 'added', p, 'added at earlier traced path; link to current path not verified'
    rc = rename_chain(rec['repo'], rec['path'], rec['rev'], offline)
    if rc and rc[0] == rec['cre']:
        f = next(f for f in files if f.get('filename') == rc[1])
        return f.get('status'), rc[1], 'added at earlier path (rename chain: ' + '; '.join(rc[2]) + ')'
    if rc:
        return '', '', f'rename chain ends at a different adding commit {rc[0][:12]} ({rc[1]})'
    base = rec['path'].rsplit('/', 1)[-1]
    same = [f for f in files if f.get('filename', '').rsplit('/', 1)[-1] == base and f.get('status') == 'added']
    if same:
        return 'added', same[0]['filename'], 'added under a different directory; needs human review'
    if c.get('_files_incomplete'):
        return '', '', 'file list incomplete: ' + c['_files_incomplete']
    return '', '', 'ADR file not found in creation commit'


def verified(fres):
    return fres.startswith('added at') and 'not verified' not in fres and 'needs human review' not in fres


def ts(v):
    return v.strftime('%Y-%m-%dT%H:%M:%SZ') if hasattr(v, 'strftime') else str(v or '')


def run(offline):
    OUT.mkdir(parents=True, exist_ok=True)
    CACHE.mkdir(parents=True, exist_ok=True)
    recs = records()
    trace, sep, jul = old_clues()
    rows, reviewed_check = [], {}
    for rec in recs:
        notes = []
        c, src = commit(rec['repo'], rec['cre'], offline)
        if '_error' in c:
            status = matched = ''
            fres = 'creation commit not retrieved: ' + c['_error']
            hits, ment, auto, rnote = [], [], 'unresolved', ''
            api_date = ''
        else:
            status, matched, fres = file_check(c, rec, trace, offline)
            hits, ment = signatures(c), mentions_only(c)
            auto, rnote = judge(hits)
            api_date = ((c.get('commit') or {}).get('committer') or {}).get('date', '')
            if api_date and api_date != ts(rec['created']):
                notes.append(f'API committer date {api_date} differs from workbook created time')
            if not verified(fres):
                auto = 'unresolved'
                notes.append('creation commit does not verifiably add this ADR; attribution shown is for the commit only')
        if rnote:
            notes.append(rnote)
        if ment:
            notes.append('tool named in message without a signature (not counted): ' + ', '.join(ment))
        if 'same commit' in rec['note']:
            notes.append('workbook note says creation and reviewed commit are the same commit, but the SHAs differ '
                         '(see suggested_corrections.csv)')
        # The reviewed commit is a confirmed record; re-detect it as a check of the rule.
        if rec['rev'] not in reviewed_check:
            rc, _ = commit(rec['repo'], rec['rev'], offline)
            reviewed_check[rec['rev']] = 'not retrieved' if '_error' in rc else \
                ('detected' if signatures(rc) else 'NOT detected')
        trailers = sum(h[0] == 'trailer' for h in hits)
        if trailers > 1:
            notes.append(f'{trailers} AI trailer lines in one message (squash-style commit)')
        s = sep.get((rec['repo'], rec['path']))
        j = jul.get((rec['repo'], rec['path']))
        j_sha = str((j or {}).get('ADR creation commit') or '')
        s_sha = str((s or {}).get('Earlier AI create/touch SHA') or '')
        ai = auto.startswith('AI attribution')
        none = auto.startswith('no qualifying')
        if j_sha and rec['cre'].startswith(j_sha[:7]) and (
                (ai and j.get('Lifecycle tag') == 'later only') or
                (none and j.get('Lifecycle tag') in ('AI at creation only', 'creation + later AI'))):
            notes.append(f"differs from 2026-07-31 table ({j.get('Lifecycle tag')}) for the same creation SHA")
        if s_sha and rec['cre'].startswith(s_sha[:7]) and none and \
                str(s.get('Lifecycle classification', '')).startswith('ADR created earlier by AI'):
            notes.append('differs from 2026-09-02 table (created earlier by AI) for the same creation SHA')
        rows.append({
            'Manual Review row': rec['row'], 'Repository': rec['repo'], 'ADR file path': rec['path'],
            'Current level': rec['level'],
            'Creation commit SHA': rec['cre'], 'Creation commit link': f"https://github.com/{rec['repo']}/commit/{rec['cre']}",
            'Reviewed commit SHA': rec['rev'], 'Reviewed commit link': f"https://github.com/{rec['repo']}/commit/{rec['rev']}",
            'Created time (UTC, workbook)': ts(rec['created']), 'Reviewed time (UTC, workbook)': ts(rec['reviewed']),
            'Creation commit committer date (API)': api_date,
            'Creation file status': status, 'Creation file matched path': matched, 'Creation file check': fres,
            'Creation commit AI attribution (automatic)': auto,
            'Evidence type(s)': '; '.join(h[0] for h in hits),
            'Evidence text': ' | '.join(h[1] for h in hits),
            'Tool / agent': ', '.join(sorted({h[2] for h in hits})),
            'Evidence source': src,
            'Reviewed commit signature re-detected': reviewed_check[rec['rev']],
            'Human review status': 'not reviewed',
            'Old clue: 2026-09-02 lifecycle': (s or {}).get('Lifecycle classification', '') or '',
            'Old clue: 2026-09-02 earlier AI SHA matches creation': ('yes' if s_sha and rec['cre'].startswith(s_sha[:7])
                                                                     else ('no' if s_sha else '')),
            'Old clue: 2026-07-31 tag': (j or {}).get('Lifecycle tag', '') or '',
            'Old clue: 2026-07-31 creation SHA matches': ('yes' if j_sha and rec['cre'].startswith(j_sha[:7])
                                                         else ('no' if j_sha else '')),
            'Notes / uncertainty': '; '.join(notes),
        })
    write(rows, recs)
    return rows


def write(rows, recs):
    with open(OUT / 'creation_ai_evidence.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    J = 'Creation commit AI attribution (automatic)'
    summ = [('scope', 'all', 'records', len(rows))]
    for lv in ('L2', 'L1', 'all'):
        sub = [r for r in rows if lv == 'all' or r['Current level'].startswith(lv)]
        for k, v in sorted(Counter(r[J] for r in sub).items()):
            summ.append(('creation commit attribution', lv, k, v))
    summ.append(('rule check', 'all', 'reviewed commits re-detected',
                 sum(v == 'detected' for v in {r['Reviewed commit SHA']: r['Reviewed commit signature re-detected'] for r in rows}.values())))
    summ.append(('rule check', 'all', 'distinct reviewed commits', len({r['Reviewed commit SHA'] for r in rows})))
    summ.append(('scope', 'all', 'distinct creation commits', len({r['Creation commit SHA'] for r in rows})))
    summ.append(('human review', 'all', 'not reviewed', sum(r['Human review status'] == 'not reviewed' for r in rows)))
    with open(OUT / 'creation_ai_summary.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['section', 'level', 'item', 'count'])
        w.writerows(summ)
    corr = [r for r in recs if 'same commit' in r['note']]
    with open(OUT / 'suggested_corrections.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['Workbook', 'Sheet', 'Manual Review row', 'Current text', 'Problem', 'Suggested text'])
        for r in corr:
            w.writerow(['adr_dataset.xlsx (same text in adr_final.xlsx)', 'Non-L3 Creation Evidence', r['row'], r['note'],
                        f"Creation commit {r['cre'][:12]} and reviewed commit {r['rev'][:12]} are different SHAs; "
                        'the reviewed commit modifies the existing ADR (verify_records_20260910.csv: fstatus = modified)',
                        'Creation and reviewed commits are different commits with identical committer timestamps.'])
        if corr:
            w.writerow(['adr_dataset.xlsx (same text in adr_final.xlsx)', 'Non-L3 Creation Summary', '',
                        'Classification review needed: ... their selected commit adds the ADR ...',
                        'Same issue as the two rows above', 'Rows ' + ' and '.join(str(r['row']) for r in corr) +
                        ' have zero intervals because the two different commits share a committer timestamp.'])


# ------------------------------------------------------------------ tests --
def selftest():
    mk = lambda msg, an='Dev', ae='dev@x.org', cn='Dev', ce='dev@x.org': {
        'commit': {'message': msg, 'author': {'name': an, 'email': ae}, 'committer': {'name': cn, 'email': ce}}}
    cases = [
        (mk('feat\n\nCo-authored-by: Claude <noreply@anthropic.com>'), 'AI attribution (automatic candidate)'),
        (mk('docs', an='Cursor Agent', ae='cursoragent@cursor.com'), 'AI attribution (automatic candidate)'),
        (mk('add\n\n🤖 Generated with [Claude Code](https://claude.com/claude-code)'), 'AI attribution (automatic candidate)'),
        (mk('fix typo', cn='copilot-swe-agent[bot]', ce='198982749+Copilot@users.noreply.github.com'),
         'AI attribution (automatic candidate)'),
        (mk('ask Claude about this later'), 'no qualifying AI attribution found'),          # mention only
        (mk('bump', an='github-actions[bot]', ae='41898282+github-actions[bot]@users.noreply.github.com'),
         'no qualifying AI attribution found'),                                            # generic bot
        (mk('x\n\nCo-authored-by: Yu Ishikawa <yu@x.org>'), 'no qualifying AI attribution found'),
        (mk('x\n\nCo-authored-by: coderabbitai[bot] <136622811+coderabbitai[bot]@users.noreply.github.com>'),
         'needs human review'),                                                             # extended list only
        (mk('x', an='Camp Dev', ae='dev@camp'), 'needs human review'),                      # loose amp\b match
    ]
    bad = [(c['commit'], want, judge(signatures(c))[0]) for c, want in cases if judge(signatures(c))[0] != want]
    assert not mentions_only(cases[0][0]) and mentions_only(cases[4][0]) == ['claude']
    for b in bad:
        print('FAIL', b)
    print(f'selftest: {len(cases) - len(bad)}/{len(cases)} rule cases pass')
    return not bad


def verify(rows):
    """Coverage and consistency checks on the written output."""
    recs = records()
    keys = [(r['Manual Review row'], r['Repository'], r['ADR file path']) for r in rows]
    lv = Counter(r['Current level'][:2] for r in rows)
    checks = {
        '71 records': len(rows) == 71,
        'no duplicate rows': len(set(keys)) == len(keys),
        'same rows as workbook': sorted(keys) == sorted((r['row'], r['repo'], r['path']) for r in recs),
        '51 L2 / 20 L1': (lv['L2'], lv['L1']) == (51, 20),
        'creation SHA never used as reviewed SHA': all(r['Creation commit SHA'] != r['Reviewed commit SHA'] for r in rows),
        'every reviewed commit re-detected as AI-signed': all(r['Reviewed commit signature re-detected'] == 'detected' for r in rows),
        'no result marked human-reviewed': all(r['Human review status'] == 'not reviewed' for r in rows),
        'unverified file addition never counted': all(r['Creation commit AI attribution (automatic)'] == 'unresolved'
                                                      for r in rows if not verified(r['Creation file check'])),
    }
    for k, v in checks.items():
        print(('ok   ' if v else 'FAIL ') + k)
    return all(checks.values())


if __name__ == '__main__':
    if '--selftest' in sys.argv:
        sys.exit(0 if selftest() else 1)
    ok = selftest()
    rows = run('--offline' in sys.argv)
    ok = verify(rows) and ok
    print(Counter((r['Current level'][:2], r['Creation commit AI attribution (automatic)']) for r in rows))
    print('output:', OUT)
    sys.exit(0 if ok else 1)
