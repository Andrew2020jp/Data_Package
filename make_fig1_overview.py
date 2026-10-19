"""Redraw ~/overview.pdf with the agreed corrections (Signal 2 = flag only,
repository triage steps, L0 exclusion step, post-L0 final dataset).
Numbers are checked against data/adr_dataset.xlsx before drawing."""
import collections, html, sys, textwrap
import openpyxl
from pathlib import Path

# ------------------------------------------------------------ numbers --
wb = openpyxl.load_workbook(str(__import__('pathlib').Path(__file__).resolve().parents[1] / 'data/adr_dataset.xlsx'), read_only=True)
A = list(wb['All Repository Checks (849)'].iter_rows(values_only=True))[1:]
M = [r for r in list(wb['Manual Review'].iter_rows(values_only=True))[6:] if r[1]]
y = lambda v: v == 'Yes'
sig = [r for r in A if y(r[10]) or y(r[13]) or y(r[15])]
body_only_out = [r for r in sig if not y(r[10]) and r[1] not in {m[1] for m in M}]
K = [r for r in M if not str(r[5]).startswith('L0')]
L0 = [r for r in M if str(r[5]).startswith('L0')]
lab = collections.Counter(str(r[7]).replace(' ', '') for r in K)
N = dict(all=len(A), sig=len(sig), body=len(body_only_out),
         cand=len(sig) - len(body_only_out), mr=len({r[1] for r in M}),
         l3=sum(str(r[5]).startswith('L3') for r in M), l2=sum(str(r[5]).startswith('L2') for r in M),
         l1=sum(str(r[5]).startswith('L1') for r in M), l0=len(L0), l0r=len({r[1] for r in L0}),
         repos=len({r[1] for r in K}), recs=len(K), commits=len({r[6] for r in K if r[6]}),
         created=lab['checked(created)'], minor=lab['checked(changed-minor)'],
         major=lab['checked(changed-major)'], move=lab['checked(move)'], rename=lab['checked(rename)'])
N['drop10'] = N['cand'] - N['mr']
EXPECT = dict(all=849, sig=191, body=74, cand=117, mr=107, drop10=10, l3=182, l2=51, l1=20,
              l0=46, l0r=12, repos=95, recs=253, commits=95, created=182, minor=34, major=17,
              move=17, rename=3)
assert N == {**N, **EXPECT} and all(N[k] == v for k, v in EXPECT.items()), N

# ------------------------------------------------------------ drawing --
FONT = 'DejaVu Sans'
C = dict(
    data=('#FDF6D8', '#E3C66B'), screen=('#FDEEDD', '#EDB27C'), sig=('#FFF6EE', '#F2C9A4'),
    ana=('#E4F4FA', '#7CC3DE'), rq=('#F1EAFB', '#A68AD6'), excl=('#F3F3F3', '#9A9A9A'),
)
INK, SUB, RED, ORG = '#1F2328', '#4A5058', '#C0392B', '#D35400'
out = []

def box(x, y, w, h, kind, dash=False, r=8):
    f, s = C[kind]
    d = ' stroke-dasharray="6 4"' if dash else ''
    out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" fill="{f}" stroke="{s}" stroke-width="1.4"{d}/>')

def text(x, y, s, size=12, weight='normal', color=INK, anchor='start', style='normal'):
    out.append(f'<text x="{x}" y="{y}" font-family="{FONT}" font-size="{size}" font-weight="{weight}" '
               f'font-style="{style}" fill="{color}" text-anchor="{anchor}">{html.escape(s)}</text>')

def para(x, y, w, s, size=12, weight='normal', color=INK, bullet=False, lh=1.32, style='normal'):
    """Wrap s into width w; returns y after the last line."""
    chars = max(10, int(w / (0.615 * size)))
    lines = textwrap.wrap(s, chars - (2 if bullet else 0))
    for i, ln in enumerate(lines):
        if bullet and i == 0:
            text(x, y, '•', size, color=color)
        text(x + (11 if bullet else 0), y, ln, size, weight, color, style=style)
        y += size * lh
    return y

def arrow(x1, y1, x2, y2, dash=False, color='#333'):
    d = ' stroke-dasharray="5 4"' if dash else ''
    out.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{color}" stroke-width="1.5"{d} marker-end="url(#ah)"/>')

def poly(pts, dash=False, color='#333', head=True):
    d = ' stroke-dasharray="5 4"' if dash else ''
    m = ' marker-end="url(#ah)"' if head else ''
    p = ' '.join(f'{a},{b}' for a, b in pts)
    out.append(f'<polyline points="{p}" fill="none" stroke="{color}" stroke-width="1.5"{d}{m}/>')

W, H = 1760, 718
TOP = 56

# legend
lx = 300
for kind, name in [('data', 'Dataset / counts'), ('screen', 'Screening criteria'),
                   ('ana', 'Analysis & manual step'), ('excl', 'Excluded'), ('rq', 'Research questions')]:
    box(lx, 16, 22, 16, kind, dash=(kind == 'excl'), r=3)
    text(lx + 30, 29, name, 12.5, color=SUB)
    lx += 60 + len(name) * 7.4

# ---- column 1: source -------------------------------------------------
x1, w1 = 12, 262
box(x1, TOP, w1, 96, 'data')
text(x1 + 14, TOP + 24, 'Dataset Source', 15, 'bold')
text(x1 + 14, TOP + 44, 'GitHub repositories collected', 12, color=SUB)
text(x1 + 14, TOP + 60, 'with Claude Code (2026)', 12, color=SUB)
text(x1 + 14, TOP + 84, f'n = {N["all"]} repositories', 15, 'bold')
arrow(x1 + w1 / 2, TOP + 96, x1 + w1 / 2, TOP + 122)

yb = TOP + 124
box(x1, yb, w1, 150, 'ana')
text(x1 + 14, yb + 24, 'Claude Code repository sweep', 13.5, 'bold')
yy = yb + 46
for s in ['locate ADR directories and files', 'record repository, file path and commit links',
          'read the ADR body text of each file']:
    yy = para(x1 + 16, yy, w1 - 30, s, 12, bullet=True)
arrow(x1 + w1 / 2, yb + 150, x1 + w1 / 2, yb + 166)

yc = yb + 168
box(x1, yc, w1, 112, 'data')
text(x1 + 14, yc + 24, 'ADR files detected', 15, 'bold')
text(x1 + 14, yc + 48, f'n = {N["all"]} repositories', 15, 'bold')
text(x1 + 14, yc + 68, 'with at least one ADR file', 12, color=SUB)
text(x1 + 14, yc + 92, '(32,808 candidate ADR files)', 12, color=SUB)

# ---- column 2: screening ----------------------------------------------
x2, w2 = 296, 300
box(x2, TOP, w2, 292, 'screen')
text(x2 + 14, TOP + 24, 'AI-trace screening', 15, 'bold')
text(x2 + 14, TOP + 42, '(Claude Code + manual verification)', 12, color=SUB)
poly([(x1 + w1, yc + 56), (x1 + w1 + 11, yc + 56), (x1 + w1 + 11, TOP + 140), (x2, TOP + 140)])

s1y = TOP + 56
box(x2 + 10, s1y, w2 - 20, 166, 'sig')
text(x2 + 22, s1y + 22, 'Signal 1 — Commit metadata', 13.5, 'bold', ORG)
text(x2 + 22, s1y + 39, '(ANY one is a hit; selects candidates)', 12, 'bold', ORG)
yy = s1y + 62
yy = para(x2 + 24, yy, w2 - 50, 'AI-related author / committer / co-author', 12, bullet=True)
yy = para(x2 + 35, yy, w2 - 60, '(e.g., Claude Code, Cursor, Copilot, Gemini, ChatGPT, etc.)', 11, color=SUB)
yy = para(x2 + 24, yy + 6, w2 - 50, 'Generated with Claude Code / Generated by ...', 12, bullet=True)

para(x2 + 14, s1y + 166 + 24, w2 - 24, 'AI wording in ADR text is not used for selection; Signal 2 is checked only for L0 records.', 10.5, color=RED, style='italic')

# ---- column 3: repository-level triage ----------------------------------
x3, w3 = 618, 236
text(x3 + w3 / 2, TOP + 14, 'Repository-level triage', 14, 'bold', anchor='middle')
arrow(x2 + w2, TOP + 60, x3, TOP + 60)

def count_box(yv, big, small, h=54):
    box(x3, yv, w3, h, 'data')
    text(x3 + w3 / 2, yv + 23, big, 14.5, 'bold', anchor='middle')
    text(x3 + w3 / 2, yv + 42, small, 11.5, color=SUB, anchor='middle')
    return yv + h

def excl_box(yv, big, lines):
    xe = x3 + 34
    h = 20 + 15 * len(lines) + 8
    box(xe, yv, w3 - 34, h, 'excl', dash=True)
    text(xe + 10, yv + 19, big, 12.5, 'bold', SUB)
    for i, ln in enumerate(lines):
        text(xe + 10, yv + 36 + 15 * i, ln, 11, color=SUB)
    poly([(x3 + 17, yv + h / 2), (xe, yv + h / 2)], head=True, color='#888')
    return yv + h

yv = count_box(TOP + 33, f'{N["sig"]} repositories', f'with any AI signal (of {N["all"]})')
out.append(f'<line x1="{x3+17}" y1="{yv}" x2="{x3+17}" y2="{yv+102}" stroke="#333" stroke-width="1.5"/>')
e = excl_box(yv + 12, f'{N["body"]} excluded', ['AI only in ADR text;', 'no commit-level AI signature'])
arrow(x3 + 17, yv + 102, x3 + 17, yv + 118)
yv = count_box(yv + 120, f'{N["cand"]} candidates', 'for manual curation')
out.append(f'<line x1="{x3+17}" y1="{yv}" x2="{x3+17}" y2="{yv+86}" stroke="#333" stroke-width="1.5"/>')
excl_box(yv + 12, f'{N["drop10"]} excluded', ['not an ADR, or link 404'])
arrow(x3 + 17, yv + 86, x3 + 17, yv + 102)
yv = count_box(yv + 104, f'{N["mr"]} repositories', 'selected for manual curation', h=58)
triage_out_y = yv - 29

# ---- column 4: manual steps ---------------------------------------------
x4, w4 = 876, 280
poly([(x3 + w3, triage_out_y), (x3 + w3 + 11, triage_out_y), (x3 + w3 + 11, TOP + 50), (x4, TOP + 50)])
box(x4, TOP, w4, 176, 'ana')
text(x4 + 14, TOP + 24, 'Manual curation and alignment', 13.5, 'bold')
yy = TOP + 46
for s in ['verify ADR body, file path and direct file link', 'align each ADR with its corresponding commit',
          'check file history / reference copies', 'apply the inclusion scope per file', 'one row per ADR file']:
    yy = para(x4 + 16, yy, w4 - 30, s, 12, bullet=True)
arrow(x4 + w4 / 2, TOP + 176, x4 + w4 / 2, TOP + 192)

yi = TOP + 194
box(x4, yi, w4, 150, 'data')
text(x4 + 14, yi + 24, 'Inclusion scope', 13.5, 'bold')
text(x4 + 14, yi + 40, '(kept as ADRs)', 11.5, color=SUB)
yy = yi + 60
for s in ['numbered ADR files', 'demo / evaluation / history / reference copies kept',
          'judged by filename or template', 'exclude non-ADR or unrelated files']:
    yy = para(x4 + 16, yy, w4 - 30, s, 12, bullet=True)
arrow(x4 + w4 / 2, yi + 150, x4 + w4 / 2, yi + 166)

yp = yi + 168
box(x4, yp, w4, 150, 'ana')
text(x4 + 14, yp + 24, 'Per-ADR commit inspection', 13.5, 'bold')
yy = yp + 46
for s in ['open the linked commit and read its metadata', 'compare the ADR file with the parent commit',
          'assign one commit-level label per ADR record']:
    yy = para(x4 + 16, yy, w4 - 30, s, 12, bullet=True)

# ---- column 5: classification / exclusion / final -----------------------
x5, w5 = 1178, 290
poly([(x4 + w4, yp + 56), (x4 + w4 + 11, yp + 56), (x4 + w4 + 11, TOP + 70), (x5, TOP + 70)])
box(x5, TOP, w5, 196, 'data')
text(x5 + 14, TOP + 24, 'Commit-level classification', 14, 'bold')
text(x5 + 14, TOP + 41, '(per ADR record, not per repository)', 11.5, color=SUB)
rows = [(f'L3 created — {N["l3"]}', 'the commit adds the ADR file', '#2E8B57'),
        (f'L2 modified — {N["l2"]}', 'existing ADR content is edited', '#1F5FAD'),
        (f'L1 rename / move only — {N["l1"]}', 'path changes, body unchanged', '#C0562B'),
        (f'L0 AI evidence unconfirmed — {N["l0"]}', 'no AI signature in the linked commit', '#6B6B6B')]
yy = TOP + 66
for big, small, col in rows:
    text(x5 + 14, yy, big, 13, 'bold', col)
    text(x5 + 14, yy + 16, small, 11.5, color=SUB)
    yy += 33

ye = TOP + 220
XE = x5 + 34
box(XE, ye, w5 - 34, 226, 'excl', dash=True)
text(XE + 12, ye + 22, f'Excluded as unconfirmed', 13, 'bold', SUB)
text(XE + 12, ye + 40, f'{N["l0"]} L0 records / {N["l0r"]} repositories', 13, 'bold', SUB)
yy = para(XE + 14, ye + 62, w5 - 64, 'no AI signature in the linked commit', 11, color=SUB, bullet=True)
box(XE + 8, yy - 6, w5 - 50, 82, 'sig', r=6)
text(XE + 16, yy + 12, 'Signal 2 — ADR body text', 12, 'bold', ORG)
yy = para(XE + 18, yy + 30, w5 - 72, 'no explicit AI authorship statement (e.g., "This ADR was written/generated by AI")', 10.5, color=SUB, bullet=True)
para(XE + 12, yy + 18, w5 - 60, 'Unconfirmed ≠ no AI use: some tools do not sign commits, so all counts are lower bounds.', 11, color=RED, style='italic')
# L0 row -> exclusion box
poly([(x5 + w5, TOP + 66 + 33 * 3 + 4), (x5 + w5 + 10, TOP + 66 + 33 * 3 + 4),
      (x5 + w5 + 10, ye + 30), (x5 + w5, ye + 30)], color='#888')

yf = ye + 246
arrow(x5 + 17, TOP + 196, x5 + 17, yf - 2)
box(x5, yf, w5, 182, 'ana')
text(x5 + 14, yf + 24, 'Final dataset (L1–L3)', 14, 'bold')
text(x5 + 14, yf + 46, f'{N["repos"]} repositories', 13.5, 'bold')
text(x5 + 14, yf + 64, f'{N["recs"]} ADR records', 13.5, 'bold')
text(x5 + 14, yf + 82, f'{N["commits"]} selected commits', 13.5, 'bold')
out.append(f'<line x1="{x5+12}" y1="{yf+94}" x2="{x5+w5-12}" y2="{yf+94}" stroke="#7CC3DE"/>')
text(x5 + 14, yf + 113, f'Human review labels ({N["recs"]} records)', 12.5, 'bold')
yy = yf + 132
LABELS = [('Created', N['created']), ('Minor modification', N['minor']),
          ('Major modification', N['major']), ('Move', N['move']), ('Rename', N['rename'])]
for i, (k, v) in enumerate(LABELS):                       # two columns to save height
    text(x5 + 22 + (0 if i < 3 else 160), yy + 16 * (i % 3), f'• {k}: {v}', 12)

# ---- column 6: research questions --------------------------------------
x6, w6 = 1494, 254
text(x6 + 14, TOP + 2, '', 1)
box(x6, TOP, w6, 470, 'rq')
text(x6 + 14, TOP + 26, 'Research Questions', 15, 'bold')
RQ = [('RQ1 — Confirmation of AI traces', 'How often are candidate AI traces supported by explicit attribution in the reviewed ADR records and selected commits?'),
      ('RQ2 - When?', 'Does AI attribution appear at file introduction or later changes? What are the creation-to-commit intervals?'),
      ('RQ3 — What kind of change?', 'Does AI write or change the actual decision content, or does it mainly make small edits such as formatting and minor fixes?')]
yy = TOP + 44
for t, d in RQ:
    by = yy
    box(x6 + 10, by, w6 - 20, 128, 'rq')
    out[-1] = out[-1].replace(C['rq'][0], '#FFFFFF')
    ty = para(x6 + 22, by + 22, w6 - 44, t, 13, 'bold')
    para(x6 + 22, ty + 3, w6 - 44, d, 11.5, color=SUB)
    poly([(x5 + w5 + 10, TOP + 30), (x5 + w5 + 14, TOP + 30), (x5 + w5 + 14, by + 56), (x6 + 10, by + 56)], dash=True, color='#666')
    yy += 138

svg = (f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">'
       '<defs><marker id="ah" markerWidth="9" markerHeight="9" refX="8" refY="4.5" orient="auto">'
       '<path d="M0,0 L9,4.5 L0,9 z" fill="#333"/></marker></defs>'
       f'<rect width="{W}" height="{H}" fill="white"/>' + ''.join(out) + '</svg>')
base = sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).resolve().parents[1] / 'figures/fig1_overview')
Path(base).parent.mkdir(exist_ok=True)
open(base + '.svg', 'w').write(svg)
import cairosvg
cairosvg.svg2pdf(bytestring=svg.encode(), write_to=base + '.pdf')
cairosvg.svg2png(bytestring=svg.encode(), write_to=base + '.png', output_width=3520)
print(N)
