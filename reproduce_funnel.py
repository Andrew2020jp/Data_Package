"""Re-run the screening steps that start from the fixed list of 849 repositories.

This script does NOT search GitHub for repositories. It starts from
`data/repositories_849.csv` (the 849 repositories collected on 2026-06-29)
and reproduces the later steps.

Two modes:

  offline  Recompute the selection funnel (849 -> 191 -> 74 -> 117 -> 10 -> 107
           -> 46 L0 -> 253 / 95) from data/adr_dataset.xlsx. No network. The numbers
           must match the paper; the script stops with an error if they do not.

  live     Query GitHub again for each repository (needs an authenticated `gh`):
             1. list ADR files on the default branch
             2. read the commits that touched the ADR directories and check
                Signal 1 (AI signature in commit metadata)
             3. read up to 30 ADR files and check Signal 2 (AI wording in text)
           Results go to output/live_screening.tsv, followed by a comparison
           with the stored screening results in data/adr_dataset.xlsx. Repositories change over time, so
           live results are not expected to match exactly.

Usage:
  python scripts/reproduce_funnel.py offline
  python scripts/reproduce_funnel.py live --limit 20
  python scripts/reproduce_funnel.py live            # all 849, slow
"""
import argparse, collections, csv, json, re, subprocess, sys, time, urllib.parse, urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]               # package root
REPOS_CSV = ROOT / 'data/repositories_849.csv'
FINAL_XLSX = ROOT / 'data/adr_dataset.xlsx'
OUT = ROOT / 'output'
CACHE = OUT / 'cache'

# ------------------------------------------------------------------ rules --
TOOLS = r'(claude|anthropic|copilot|cursor|gemini|chatgpt|openai|codex|devin|aider|windsurf)'
TRAILER = re.compile(r'^\s*co-authored-by:.*' + TOOLS, re.I | re.M)
WATERMARK = re.compile(r'generated (with|by)\s+\[?' + TOOLS, re.I)
IDENTITY = re.compile(TOOLS, re.I)
KEYWORD = re.compile(TOOLS, re.I)
BROAD = re.compile(r'\b(Claude|Cursor|ChatGPT|OpenAI|LLM|Copilot|Gemini)\b')
STRICT = [
    re.compile(r'\b(this|the)\s+(adr|document|record|decision record)\s+(was|is|has been)\s+'
               r'(written|generated|drafted|authored|created)\s+(by|with)\s+(an?\s+)?'
               r'(ai|llm|claude|chatgpt|copilot|gpt[-\d.]*|gemini)\b', re.I),
    re.compile(r'\bai[- ]generated\s+(adr|architecture decision)', re.I),
    re.compile(r'\b(adr|architecture decision record)s?\s+(written|generated)\s+by\s+(ai|an llm|claude|chatgpt)\b', re.I),
]
ADR_PATH = re.compile(r'(^|/)(adr|adrs|decisions|decision-records?|architecture[-_]decisions?'
                      r'|architecture-decision-records?|architecture-decision-log|decision-log|architecture/decisions)/', re.I)
NOT_ADR = re.compile(r'(^|/)(readme|index|template|_template|adr-template)[^/]*\.md$', re.I)
SCAN_CAP = 30                                            # same cap as the original scan


# ---------------------------------------------------------------- offline --
def offline():
    import openpyxl
    wb = openpyxl.load_workbook(FINAL_XLSX, read_only=True)
    A = list(wb['All Repository Checks (849)'].iter_rows(values_only=True))[1:]
    M = [r for r in list(wb['Manual Review'].iter_rows(values_only=True))[6:] if r[1]]
    yes = lambda v: v == 'Yes'
    mr = {r[1] for r in M}
    sig = [r for r in A if yes(r[10]) or yes(r[13]) or yes(r[15])]
    body_only_out = [r for r in sig if not yes(r[10]) and r[1] not in mr]
    cand = len(sig) - len(body_only_out)
    l0 = [r for r in M if str(r[5]).startswith('L0')]
    kept = [r for r in M if not str(r[5]).startswith('L0')]
    got = {
        'repositories collected': len(A),
        'with any AI signal': len(sig),
        'excluded: AI only in ADR text': len(body_only_out),
        'candidates': cand,
        'excluded: non-ADR or 404': cand - len(mr),
        'manual review repositories': len(mr),
        'manual review records': len(M),
        'manual review commits': len({r[6] for r in M}),
        'L0 records': len(l0),
        'L0 repositories': len({r[1] for r in l0}),
        'final records': len(kept),
        'final repositories': len({r[1] for r in kept}),
        'final commits': len({r[6] for r in kept}),
    }
    want = [849, 191, 74, 117, 10, 107, 299, 108, 46, 12, 253, 95, 95]
    for (k, v), w in zip(got.items(), want):
        print(f'{k:32s} {v:5d}' + ('' if v == w else f'   <-- expected {w}'))
    if list(got.values()) != want:
        sys.exit('MISMATCH: funnel does not match the paper')
    print('\nOK: all funnel numbers match the paper.')


# ------------------------------------------------------------------- live --
def gh(path, cache_key):
    f = CACHE / (re.sub(r'[^A-Za-z0-9_.-]', '_', cache_key) + '.json')
    if f.exists():
        return json.loads(f.read_text())
    for attempt in range(3):
        p = subprocess.run(['gh', 'api', '-H', 'Accept: application/vnd.github+json', path],
                           capture_output=True, text=True)
        if p.returncode == 0:
            data = json.loads(p.stdout or 'null')
            f.write_text(json.dumps(data))
            return data
        if 'rate limit' in p.stderr.lower():
            time.sleep(60)
            continue
        data = {'_error': p.stderr.strip()[:200]}
        f.write_text(json.dumps(data))
        return data
    return {'_error': 'rate limited'}


def raw(repo, branch, path):
    f = CACHE / ('raw_' + re.sub(r'[^A-Za-z0-9_.-]', '_', f'{repo}@{branch}/{path}'))
    if f.exists():
        return f.read_text(errors='ignore')
    url = f'https://raw.githubusercontent.com/{repo}/{branch}/{urllib.parse.quote(path)}'
    try:
        with urllib.request.urlopen(url, timeout=20) as r:
            t = r.read().decode('utf-8', errors='ignore')
    except Exception:
        t = ''
    f.write_text(t)
    return t


def check_commit(c):
    msg = (c.get('commit') or {}).get('message', '') or ''
    who = ' '.join(filter(None, [
        (c.get('commit') or {}).get('author', {}).get('name'),
        (c.get('commit') or {}).get('author', {}).get('email'),
        (c.get('commit') or {}).get('committer', {}).get('name'),
        (c.get('author') or {}).get('login') if isinstance(c.get('author'), dict) else None,
    ]))
    hits = []
    if TRAILER.search(msg):
        hits.append('trailer')
    if WATERMARK.search(msg):
        hits.append('watermark')
    if IDENTITY.search(who):
        hits.append('identity')
    tools = sorted({m.lower() for m in re.findall(TOOLS, msg + ' ' + who, re.I)})
    return hits, bool(KEYWORD.search(msg)), tools


def screen(repo):
    info = gh(f'repos/{repo}', f'repo_{repo}')
    if '_error' in info:
        return dict(repo=repo, status='repo error: ' + info['_error'][:60])
    branch = info['default_branch']
    tree = gh(f'repos/{repo}/git/trees/{branch}?recursive=1', f'tree_{repo}')
    files = [t['path'] for t in (tree.get('tree') or []) if t.get('type') == 'blob'
             and t['path'].lower().endswith('.md') and ADR_PATH.search(t['path'])
             and not NOT_ADR.search(t['path'])]
    dirs = sorted({f.rsplit('/', 1)[0] for f in files})[:5]

    sig_commits, kw_commits, tools = [], [], set()
    for d in dirs:
        cs = gh(f'repos/{repo}/commits?path={urllib.parse.quote(d)}&per_page=100', f'commits_{repo}_{d}')
        for c in cs if isinstance(cs, list) else []:
            hits, kw, t = check_commit(c)
            if hits:
                sig_commits.append(c['html_url']); tools.update(t)
            elif kw:
                kw_commits.append(c['html_url'])

    broad, strict = set(), []
    for p in files[:SCAN_CAP]:
        text = raw(repo, branch, p)
        broad.update(BROAD.findall(text))
        if any(rx.search(text) for rx in STRICT):
            strict.append(p)

    return dict(repo=repo, status='ok', adr_files=len(files), adr_files_scanned=min(len(files), SCAN_CAP),
                signal1_signature='Yes' if sig_commits else 'No',
                commit_keyword_only='Yes' if kw_commits and not sig_commits else 'No',
                tools=';'.join(sorted(tools)), example_ai_commit=(sig_commits or [''])[0],
                signal2_broad='Yes' if broad else 'No', broad_words='; '.join(sorted(broad)),
                signal2_strict='Yes' if strict else 'No', strict_files=' '.join(strict[:3]))


def live(limit, only):
    CACHE.mkdir(parents=True, exist_ok=True)
    repos = [r['full_name'] for r in csv.DictReader(open(REPOS_CSV))]
    if only:
        repos = [r for r in repos if r in only]
    if limit:
        repos = repos[:limit]
    rows = []
    for i, r in enumerate(repos, 1):
        rows.append(screen(r))
        print(f'[{i}/{len(repos)}] {r}: {rows[-1].get("status")}', file=sys.stderr)
    cols = ['repo', 'status', 'adr_files', 'adr_files_scanned', 'signal1_signature', 'commit_keyword_only',
            'tools', 'example_ai_commit', 'signal2_broad', 'broad_words', 'signal2_strict', 'strict_files']
    out = OUT / 'live_screening.tsv'
    with out.open('w', newline='') as f:
        w = csv.DictWriter(f, cols, delimiter='\t', extrasaction='ignore')
        w.writeheader(); w.writerows(rows)

    ok = [r for r in rows if r['status'] == 'ok']
    s1 = sum(r['signal1_signature'] == 'Yes' for r in ok)
    kw = sum(r['commit_keyword_only'] == 'Yes' for r in ok)
    body = sum(r['signal1_signature'] == 'No' and r['commit_keyword_only'] == 'No'
               and (r['signal2_broad'] == 'Yes' or r['signal2_strict'] == 'Yes') for r in ok)
    print(f'\nLive results ({len(rows)} repositories, {len(ok)} reachable) -> {out}')
    print(f'  Signal 1, real AI signature in an ADR commit : {s1}')
    print(f'  commit message only names an AI tool         : {kw}')
    print(f'  AI only in ADR text (Signal 2, no commit)    : {body}')

    import openpyxl
    sheet = openpyxl.load_workbook(FINAL_XLSX, read_only=True)['All Repository Checks (849)']
    stored = {r[1]: r for r in list(sheet.iter_rows(values_only=True))[1:]}
    agree = collections.Counter()
    for r in ok:
        s = stored.get(r['repo'])
        if not s:
            continue
        old_commit = s[10] == 'Yes'                        # Commit AI trace
        new_commit = r['signal1_signature'] == 'Yes' or r['commit_keyword_only'] == 'Yes'
        agree['commit trace same' if old_commit == new_commit else 'commit trace differs'] += 1
        agree['broad same' if (s[13] == 'Yes') == (r['signal2_broad'] == 'Yes')   # Broad AI wording
              else 'broad differs'] += 1
    print('\nCompared with the stored screening results:')
    for k, v in sorted(agree.items()):
        print(f'  {k:22s} {v}')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('mode', choices=['offline', 'live'])
    ap.add_argument('--limit', type=int, default=0, help='live: only the first N repositories')
    ap.add_argument('--repos', nargs='*', default=[], help='live: only these repositories')
    a = ap.parse_args()
    offline() if a.mode == 'offline' else live(a.limit, set(a.repos))
