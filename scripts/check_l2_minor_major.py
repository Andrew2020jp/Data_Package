"""AI check: is each L2 change minor or major?

Scope: the 51 L2 records. For each one, the script takes the ADR file's diff in the reviewed
commit and decides minor / major by a fixed rule (below). It does not read the human labels
when deciding; it only compares with them afterwards.

  python scripts/check_l2_minor_major.py            # check, write CSV, print agreement with human check
  python scripts/check_l2_minor_major.py --write    # also write the "AI check (L2 minor/major)" column
                                                    # in Manual Review and Non-L3 Creation Evidence
                                                    # (added if missing, otherwise only its L2 cells updated)

Run from data_package/. Diffs come from the cached commit responses in
analysis/non_l3_creation_ai/api_cache/. When GitHub left the patch out (commit with >= 3,000
files), the file is fetched before and after the commit with `gh api` and diffed locally; those
files are cached in analysis/l2_minor_major/blob_cache/.

Rule (the paper's definition). Minor: formatting, spelling, punctuation, status or metadata
updates, and path or reference changes that do not alter the documented decision content.
Major: changes to the documented problem context, decision, rationale, or trade-offs.

The script applies it like this. Each changed line is placed under its heading in the full file
at the reviewed commit. Every changed word is then sorted:
  status   - under a status, progress or reference heading (Status, Last Updated, Related,
             References, Follow-ups, Future, Rollout, Progress, Changelog, History, and the
             Korean/German words for "current state": 현황, Stand), in the preamble before the
             first section, or in a block starting with Status:, Date:, Last-revised:, Deciders:;
  format   - the block is identical after removing markdown marks, punctuation, table padding
             and blank lines;
  spelling - one word replaced by one similar word (character similarity >= 0.6, e.g.
             imlemented -> implemented, ETC -> ECT);
  path/ref - a replacement of at most 3 words by at most 3 words where one side is a path,
             link, domain or file name (contains "/" or a dot between letters), or any such
             replacement inside a code sample (identifier rename);
  content  - everything else.
major = at least one content word changed; otherwise minor. Any heading not listed under
status counts as decision content, because an ADR's sections document the decision.

Version 2 (2026-10-23). Version 1 treated progress sections (Open follow-ups, Rollout, 현황) and
identifier renames in code samples as content; a review of all 51 diffs against the definition
found 6 wrong labels (rows 11, 12, 54, 58, 74, 144). Version 2 adds the progress headings and
the code-sample rule above. Row 74 (proposed -> accepted, written as prose inside the ADR) is
not caught by any rule, so it is set in OVERRIDES with the reason. Version 1 output is kept in
analysis/l2_minor_major/l2_ai_check_v1.csv.
"""
import csv
import difflib
import json
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'data/adr_dataset.xlsx'
CACHE = ROOT / 'analysis/non_l3_creation_ai/api_cache'
OUT = ROOT / 'analysis/l2_minor_major'
BLOBS = OUT / 'blob_cache'
BACKUP = ROOT.parent / 'others/backup'

COL = 'AI check (L2 minor/major)'
STATUS_HEAD = re.compile(r'^\s*#+\s*(status|last updated|related|references|open follow-ups?|follow-ups?|future'
                         r'|rollout|progress|changelog|history|stand\b|\S*\s*현황)', re.I)
STATUS_WORDS = re.compile(r'(status|rollout|follow-ups?|future|progress|현황|stand)\b', re.I)
STATUS_LINE = re.compile(r'^[\s>*\-]*\**\s*(status|date|last[- ]revised|last[- ]updated|deciders)\s*\**\s*:', re.I)
TOKEN = re.compile(r"[A-Za-z0-9]+(?:[./:_'’-]+[A-Za-z0-9]+)*|[぀-ヿ㐀-鿿가-힯]")
# review decisions the rules cannot make: Manual Review row -> (label, reason)
OVERRIDES = {
    74: ('minor', 'proposed -> accepted written as prose (Status line, preamble note, "Decision Outcome'
                  ' (recommended, proposed)" heading, 推奨 -> 採用); the chosen options are unchanged'),
}
CJK = re.compile(r'[぀-ヿ㐀-鿿가-힯]')


# ------------------------------------------------------------------ diffs --
def cached_commit(repo, sha):
    files = []
    for f in sorted(CACHE.glob(f"{repo.replace('/', '__')}__{sha}__p*.json")):
        body = json.loads(f.read_text())['body']
        files += body.get('files') or []
        parent = body['parents'][0]['sha'] if body.get('parents') else None
    return files, parent


def gh_file(repo, path, ref):
    f = BLOBS / f"{repo.replace('/', '__')}__{ref}__{path.replace('/', '__')}"
    if f.exists():
        return f.read_text()
    r = subprocess.run(['gh', 'api', '-H', 'Accept: application/vnd.github.raw',
                        f'repos/{repo}/contents/{path}?ref={ref}'], capture_output=True, text=True)
    if r.returncode:
        return None
    BLOBS.mkdir(parents=True, exist_ok=True)
    f.write_text(r.stdout)
    return r.stdout


def get_patch(repo, sha, path):
    """(patch text, source note) for the ADR file in the reviewed commit."""
    files, parent = cached_commit(repo, sha)
    hit = [f for f in files if f['filename'] == path]
    if hit and hit[0].get('patch'):
        return hit[0]['patch'], f"cached commit patch ({hit[0]['status']} +{hit[0]['additions']}/-{hit[0]['deletions']})"
    old, new = gh_file(repo, path, parent), gh_file(repo, path, sha)
    if old is None or new is None:
        return None, 'diff unavailable'
    lines = list(difflib.unified_diff(old.splitlines(), new.splitlines(), lineterm='', n=3))[2:]
    return '\n'.join(lines), f'gh api file before/after (commit lists {len(files)} files, patch omitted)'


# ------------------------------------------------------------------- rule --
def norm(line):
    s = re.sub(r'[*`~>#|]', ' ', line)
    s = re.sub(r'(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])', ' ', s)
    s = re.sub(r'-{3,}', ' ', s)
    return TOKEN.findall(s)


def weight(tokens):
    return sum(0.5 if CJK.fullmatch(t) else 1 for t in tokens)


def is_ref(t):
    return '/' in t or re.search(r'[A-Za-z]\.[A-Za-z]', t) is not None


def headings(text):
    """[(line number, level, title)] outside code fences, and the set of lines inside fences."""
    out, code, fence = [], set(), False
    for i, line in enumerate(text.splitlines(), 1):
        if line.lstrip().startswith('```'):
            fence = not fence
            continue
        if fence:
            code.add(i)
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m and not fence:
            out.append((i, len(m.group(1)), m.group(2).strip()))
    return out, code


def section(heads, pos):
    """'status' or the heading path ('' = preamble) for new-file line pos."""
    stack = {}
    for i, lv, t in heads:
        if i > pos:
            break
        stack = {k: v for k, v in stack.items() if k < lv}
        stack[lv] = t
    path = [stack[k] for k in sorted(stack)]
    if len(path) <= 1 or any(STATUS_HEAD.match('# ' + t) or STATUS_WORDS.search(t) for t in path[1:]):
        return 'status', ' > '.join(path[1:])
    return 'content', ' > '.join(path[1:])


def blocks(patch, heads, code):
    """(removed lines, added lines, section kind, section name, in code) for each change block."""
    out, rem, add, where, nl = [], [], [], None, 0
    def flush():
        if rem or add:
            out.append((rem[:], add[:]) + where)
        rem.clear(); add.clear()
    for line in patch.splitlines():
        h = re.match(r'@@ -\d+(?:,\d+)? \+(\d+)', line)
        if h:
            flush(); nl = int(h.group(1)); continue
        tag, body = line[:1], line[1:]
        if tag in ('-', '+'):
            if not rem and not add:
                where = section(heads, nl) + (nl in code,)
            (rem if tag == '-' else add).append(body)
            nl += tag == '+'
            continue
        flush(); nl += 1
    flush()
    return out


def classify(patch, text):
    c = dict(status=0, format=0, spelling=0, ref=0, content=0.0)
    where = []
    for rem, add, kind, name, in_code in blocks(patch, *headings(text)):
        first = next((l for l in rem + add if l.strip()), '')
        if kind == 'status' or STATUS_LINE.match(first):
            c['status'] += 1
            continue
        a, b = norm(' '.join(rem)), norm(' '.join(add))
        if a == b:
            c['format'] += 1
            continue
        for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
            x, y = a[i1:i2], b[j1:j2]
            if op == 'equal':
                continue
            if op == 'replace' and len(x) <= 3 and len(y) <= 3 and (in_code or any(map(is_ref, x + y))):
                c['ref'] += 1
            elif op == 'replace' and len(x) == len(y) == 1 and \
                    difflib.SequenceMatcher(None, x[0].lower(), y[0].lower()).ratio() >= 0.6:
                c['spelling'] += 1
            else:
                c['content'] += max(weight(x), weight(y))
                if name not in where:
                    where.append(name)
    label = 'major' if c['content'] else 'minor'
    return label, c, where


# ------------------------------------------------------------------ inputs --
def full_sha(cell):
    m = re.search(r'/commit/([0-9a-f]{40})', (cell.hyperlink.target if cell.hyperlink else '') or '')
    return m.group(1) if m else None


def header_row(ws, first):
    for r in ws.iter_rows(max_row=10):
        if r[0].value == first:
            return r[0].row, {c.value.strip() if isinstance(c.value, str) else c.value: c.column for c in r}
    raise SystemExit(f'header not found in {ws.title}')


def run():
    wb = openpyxl.load_workbook(BOOK)
    ws = wb['Non-L3 Creation Evidence']
    hr, col = header_row(ws, 'Manual Review row')
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        if not str(ws.cell(r, col['Commit Level']).value or '').startswith('L2'):
            continue
        repo, path = ws.cell(r, col['Repository']).value, ws.cell(r, col['ADR file path']).value
        sha = full_sha(ws.cell(r, col['Reviewed commit']))
        patch, src = get_patch(repo, sha, path)
        text = gh_file(repo, path, sha)
        if patch is None or text is None:
            label, c, where = 'unresolved', dict(status=0, format=0, spelling=0, ref=0, content=0), []
        else:
            label, c, where = classify(patch, text)
        row = ws.cell(r, col['Manual Review row']).value
        rule_label = label
        label, why = OVERRIDES.get(row, (label, ''))
        rows.append({'Manual Review row': ws.cell(r, col['Manual Review row']).value,
                     'Repository': repo, 'ADR file path': path, 'Reviewed commit SHA': sha,
                     COL: label, 'Rule result': rule_label, 'Override reason': why,
                     'Human check (L2 change)': ws.cell(r, col['L2 change (minor/major)']).value,
                     'Content words changed': round(c['content'], 1),
                     'Content changed in sections': ' | '.join(where),
                     'Spelling replacements': c['spelling'], 'Path/reference replacements': c['ref'],
                     'Format-only blocks': c['format'], 'Status/metadata blocks': c['status'],
                     'Diff source': src})
    return rows


def save_csv(rows):
    OUT.mkdir(parents=True, exist_ok=True)
    f = OUT / 'l2_ai_check.csv'
    with f.open('w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)
    return f


def report(rows):
    from collections import Counter
    ai = Counter(r[COL] for r in rows)
    hu = Counter(r['Human check (L2 change)'] for r in rows)
    both = Counter((r['Human check (L2 change)'], r[COL]) for r in rows)
    print(f'L2 records: {len(rows)}')
    print(f"Human check: minor {hu['minor']}, major {hu['major']}")
    assert (hu['minor'], hu['major']) == (37, 14), 'human labels no longer match the paper (37 minor, 14 major)'
    print(f"AI check:    minor {ai['minor']}, major {ai['major']}, unresolved {ai['unresolved']}")
    print(f"Same: {both[('minor','minor')] + both[('major','major')]}  "
          f"(human minor -> AI major: {both[('minor','major')]}, human major -> AI minor: {both[('major','minor')]})")
    for r in rows:
        if r[COL] != r['Human check (L2 change)']:
            print(f"  row {r['Manual Review row']}: human {r['Human check (L2 change)']}, AI {r[COL]}, "
                  f"content {r['Content words changed']} words  {r['Repository']} {r['ADR file path']}")


# ------------------------------------------------------------------- write --
def snapshot(wb):
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb}


def add_column(ws, hr, key_cols, labels, blank_value):
    """Append COL after the last used column; fill by (repo, path)."""
    _, col = header_row(ws, ws.cell(hr, 1).value)
    if COL in col:
        new = col[COL]
    else:
        new = ws.max_column + 1
        h = ws.cell(hr, new, COL)
        h._style = ws.cell(hr, new - 1)._style
        ws.column_dimensions[get_column_letter(new)].width = 16
    filled = 0
    for r in range(hr + 1, ws.max_row + 1):
        level = str(ws.cell(r, col['Commit Level']).value or '')
        key = tuple(ws.cell(r, col[k]).value for k in key_cols)
        if not level.startswith(('L0', 'L1', 'L2', 'L3', 'N/A')):
            continue
        cell = ws.cell(r, new)
        if cell.value is None:
            cell._style = ws.cell(r, new - 1)._style
        if level.startswith('L2'):
            cell.value = labels.pop(key)
            filled += 1
        elif blank_value:
            cell.value = blank_value
    return filled


def hr_of(ws):
    for first in ('#', 'Manual Review row'):
        try:
            return header_row(ws, first)[0]
        except SystemExit:
            pass
    return 1


def write(rows):
    BACKUP.mkdir(parents=True, exist_ok=True)
    bk = BACKUP / f'adr_dataset_before_l2_ai_check_{time.strftime("%Y%m%d_%H%M%S")}.xlsx'
    shutil.copy2(BOOK, bk)
    wb = openpyxl.load_workbook(BOOK)
    before = snapshot(wb)
    labels = {(r['Repository'], r['ADR file path']): r[COL] for r in rows}

    ws = wb['Manual Review']
    hr, _ = header_row(ws, '#')
    n1 = add_column(ws, hr, ('Repository', 'ADR file path'), dict(labels), None)
    ws = wb['Non-L3 Creation Evidence']
    hr, _ = header_row(ws, 'Manual Review row')
    n2 = add_column(ws, hr, ('Repository', 'ADR file path'), dict(labels), 'Not L2')
    assert n1 == n2 == len(rows) == 51, (n1, n2, len(rows))

    # every cell outside the AI check column unchanged
    after = snapshot(wb)
    for name, old in before.items():
        skip = {i for i, v in enumerate(after[name][hr_of(wb[name]) - 1]) if v == COL}
        for i, row in enumerate(old):
            assert all(after[name][i][j] == v for j, v in enumerate(row) if j not in skip), f'{name} row {i + 1} changed'
    wb.save(BOOK)
    print(f'written: {BOOK}  (backup {bk})')


if __name__ == '__main__':
    rows = run()
    print('CSV:', save_csv(rows))
    report(rows)
    if '--write' in sys.argv:
        write(rows)
