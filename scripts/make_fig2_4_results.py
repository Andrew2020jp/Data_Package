"""Redraw the three results figures (Fig. 2-4) for the L0-excluded version.

Numbers are read from data/adr_dataset.xlsx and checked against the paper.
Style follows the original figures: DejaVu Serif, 3.45 x 2.75 in.
Output: figures/results_rq{1,2,3}.pdf
"""
import collections, statistics
from pathlib import Path
import openpyxl
import matplotlib
matplotlib.use('Agg')
(Path(__file__).resolve().parents[1] / 'figures').mkdir(exist_ok=True)
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
wb = openpyxl.load_workbook(ROOT / 'data/adr_dataset.xlsx', read_only=True, data_only=True)
M = [r for r in list(wb['Manual Review'].iter_rows(values_only=True))[6:] if r[1]]
lv = collections.Counter(str(r[5])[:2] for r in M)
lab = collections.Counter(str(r[7]).replace(' ', '') for r in M if str(r[5]).startswith('L2'))
E = [r for r in list(wb['Non-L3 Creation Evidence'].iter_rows(values_only=True))[3:] if r[0]]
days = lambda r: (r[5] - r[4]).total_seconds() / 86400
med = {L: round(statistics.median(days(r) for r in E if str(r[2]).startswith(L)), 2) for L in ('L2', 'L1')}
n = {L: sum(str(r[2]).startswith(L) for r in E) for L in ('L2', 'L1')}
minor, major = lab['checked(changed-minor)'], lab['checked(changed-major)']
assert (len(M), lv['L3'], lv['L2'], lv['L1'], lv['L0']) == (299, 182, 51, 20, 46)
assert med == {'L2': 5.33, 'L1': 129.15} and n == {'L2': 51, 'L1': 20}
assert (minor, major) == (34, 17)

plt.rcParams.update({'font.family': 'DejaVu Serif', 'font.size': 9,
                     'axes.spines.top': False, 'axes.spines.right': False})
DARK, LIGHT, GREY = '#3f6b94', '#89a9c0', '#d0d0d0'


def base(title, note):
    fig, ax = plt.subplots(figsize=(3.45, 2.75))
    ax.set_title(title, loc='right', fontsize=8.5)
    ax.yaxis.grid(True, color='#e3e3e3', linewidth=0.6); ax.set_axisbelow(True)
    fig.text(0.5, 0.02, note, ha='center', fontsize=7)
    return fig, ax


def label(ax, bars, texts):
    for b, t in zip(bars, texts):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + ax.get_ylim()[1] * 0.02, t,
                ha='center', va='bottom', fontsize=8.5)


# Fig. 2: RQ1, all 299 reviewed records; L0 shown but marked as excluded
fig, ax = base(f'{len(M)} reviewed ADR records',
               'L0 attribution is unconfirmed; L0 is excluded from RQ2 and RQ3.')
vals = [lv['L3'], lv['L2'], lv['L1'], lv['L0']]
bars = ax.bar(range(4), vals, width=0.62, color=[DARK, DARK, DARK, GREY], edgecolor='#444', linewidth=0.6)
bars[3].set_hatch('///')
ax.set_ylim(0, 225)
label(ax, bars, [f'{v}\n({v / len(M):.1%})' for v in vals])
ax.set_xticks(range(4), ['ADR\ncreation\n(L3)', 'Content\nmodification\n(L2)', 'Rename/\nmove\n(L1)',
                         'Unconfirmed\n(L0,\nexcluded)'], fontsize=8)
ax.set_ylabel('Number of ADR records')
fig.subplots_adjust(left=0.18, right=0.98, top=0.9, bottom=0.3)
fig.savefig(ROOT / 'figures/results_rq1.pdf'); plt.close(fig)

# Fig. 3: RQ2, confirmed L1/L2 only
fig, ax = base('Creation to reviewed commit (UTC)',
               'Confirmed attribution only (L0 excluded).')
bars = ax.bar([0, 1], [med['L2'], med['L1']], width=0.52, color=[DARK, LIGHT], edgecolor='#444', linewidth=0.6)
ax.set_ylim(0, 155)
label(ax, bars, [f'{med["L2"]:.2f}', f'{med["L1"]:.2f}'])
ax.set_xticks([0, 1], [f'Modification\n(L2, n = {n["L2"]})', f'Rename/move\n(L1, n = {n["L1"]})'], fontsize=8)
ax.set_ylabel('Median interval (days)')
fig.subplots_adjust(left=0.18, right=0.98, top=0.9, bottom=0.3)
fig.savefig(ROOT / 'figures/results_rq2.pdf'); plt.close(fig)

# Fig. 4: RQ3, confirmed L2 only
tot = minor + major
fig, ax = base(f'Human review labels (L2, n = {tot})',
               'Confirmed L2 records only (L0 excluded).')
bars = ax.bar([0, 1], [minor, major], width=0.52, color=[LIGHT, DARK], edgecolor='#444', linewidth=0.6)
ax.set_ylim(0, 43)
label(ax, bars, [f'{minor}\n({minor / tot:.1%})', f'{major}\n({major / tot:.1%})'])
ax.set_xticks([0, 1], ['Minor changes', 'Major changes'], fontsize=8)
ax.set_ylabel('Number of ADR records')
fig.subplots_adjust(left=0.18, right=0.98, top=0.9, bottom=0.25)
fig.savefig(ROOT / 'figures/results_rq3.pdf'); plt.close(fig)
print('ok', dict(lv), med, minor, major)
